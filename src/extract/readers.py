"""Plain-text extraction from uploaded documents and scraped decision artefacts.

v9.633 — il FORMATO si legge dal CONTENUTO, non dall'estensione. Caso vero (12 ott 2026): l'avvocato ha caricato
«Autorizim Perdorim Auto mjeti SHPK.doc», che è un .docx (zip OOXML) rinominato .doc; il lettore andava ad antiword per
l'estensione e, per giunta, cercava antiword in /usr/local/bin mentre Debian lo installa in /usr/bin: OGNI .doc caricato finiva
in errore, e il cervello ha «rivisto» il documento senza averlo mai letto. Ora:

- ``formato_reale(path)`` guarda le prime righe del file: zip con ``word/document.xml`` → docx, zip OpenDocument → odt,
  OLE2 (D0 CF 11 E0) → doc, ``%PDF`` → pdf, ``{\\rtf`` → rtf, HTML, immagini (JPEG/PNG/WebP/HEIC);
- ``read_text`` usa il formato reale quando lo riconosce, l'estensione solo altrimenti;
- antiword si trova con ``shutil.which`` e scrive in UTF-8 (``-m UTF-8.txt``: senza, la «ë» usciva come carattere rotto);
- dal .docx si legge il corpo NELL'ORDINE (paragrafi e tabelle), le caselle di testo, le intestazioni e i piè di pagina, le
  note: prima solo i paragrafi (python-docx ``doc.paragraphs``), e le tabelle di un contratto sparivano in silenzio;
- RTF e ODT con un lettore proprio (prima l'RTF arrivava al modello con i codici di formattazione).

``read_text(path)`` returns a ``ReadResult`` so downstream callers can tell an empty document from an extraction failure.
"""
from __future__ import annotations

import logging
import os
import re
import shutil
import subprocess
import zipfile
from dataclasses import dataclass
from pathlib import Path

log = logging.getLogger(__name__)


def _antiword_bin() -> str:
    """Il binario vero: Debian lo mette in /usr/bin (il percorso fisso /usr/local/bin non esisteva nel container)."""
    trovato = shutil.which("antiword")
    if trovato:
        return trovato
    for p in ("/usr/bin/antiword", "/usr/local/bin/antiword"):
        if os.path.exists(p):
            return p
    return "antiword"


# compatibilita': qualcuno potrebbe leggerla
ANTIWORD_BIN = _antiword_bin()


@dataclass
class ReadResult:
    text: str
    source_format: str
    pages: int | None = None
    # True if extraction ran without raising; the text itself may still
    # be empty (e.g. a scanned PDF with no embedded text).
    ok: bool = True
    error: str | None = None


# ── il formato dal contenuto ──────────────────────────────────────────────────


def formato_reale(path: str | Path) -> str | None:
    """Il formato vero del file dai suoi primi byte: 'docx' | 'odt' | 'doc' | 'pdf' | 'rtf' | 'html' | 'jpg' | 'png' |
    'webp' | 'heic' | None (testo semplice o sconosciuto). Mai solleva."""
    try:
        with open(path, "rb") as fh:
            testa = fh.read(2048)
    except OSError:
        return None
    if not testa:
        return None
    if testa.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(path) as z:
                nomi = set(z.namelist())
                if "word/document.xml" in nomi:
                    return "docx"
                if "content.xml" in nomi and "mimetype" in nomi:
                    mt = z.read("mimetype")[:80]
                    if b"opendocument.text" in mt:
                        return "odt"
        except Exception:  # noqa: BLE001 — uno zip rotto non è un documento
            return None
        return None
    if testa.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "doc"          # OLE2: Word 97-2003 (o un altro file Office: antiword lo dirà)
    if b"%PDF-" in testa[:1024]:
        return "pdf"
    if testa.lstrip(b"\xef\xbb\xbf \t\r\n").startswith(b"{\\rtf"):
        return "rtf"
    if testa[:3] == b"\xff\xd8\xff":
        return "jpg"
    if testa.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if testa[:4] == b"RIFF" and testa[8:12] == b"WEBP":
        return "webp"
    if testa[4:8] == b"ftyp" and testa[8:12] in (b"heic", b"heix", b"hevc", b"heim", b"heis", b"mif1", b"msf1"):
        return "heic"
    t = testa.lstrip(b"\xef\xbb\xbf \t\r\n").lower()
    if t.startswith(b"<!doctype html") or t.startswith(b"<html") or (b"<html" in t[:600] and b"<body" in testa.lower()):
        return "html"
    return None


# ── PDF ───────────────────────────────────────────────────────────────────────


def _read_pdf(path: Path) -> ReadResult:
    import pdfplumber  # heavy import; defer
    try:
        pieces: list[str] = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                pieces.append(page.extract_text() or "")
            pages = len(pdf.pages)
        return ReadResult(
            text="\n".join(pieces).strip(),
            source_format="pdf",
            pages=pages,
        )
    except Exception as e:
        return ReadResult(text="", source_format="pdf", ok=False, error=str(e))


# ── Word 97-2003 (.doc) ───────────────────────────────────────────────────────


def _read_doc(path: Path) -> ReadResult:
    # antiword is a tiny C tool tuned for Word 97-2003 binary .doc files.
    # `-w 0` disables word wrap so we preserve paragraph structure; `-m UTF-8.txt` writes ë/ç as UTF-8.
    try:
        result = subprocess.run(
            [_antiword_bin(), "-m", "UTF-8.txt", "-w", "0", str(path)],
            capture_output=True,
            timeout=30,
            check=False,
        )
        if result.returncode != 0:
            return ReadResult(
                text="",
                source_format="doc",
                ok=False,
                error=result.stderr.decode("utf-8", "replace")[:300],
            )
        return ReadResult(
            text=result.stdout.decode("utf-8", "replace").strip(),
            source_format="doc",
        )
    except Exception as e:
        return ReadResult(text="", source_format="doc", ok=False, error=str(e))


# ── Word 2007+ (.docx) ────────────────────────────────────────────────────────

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _testo_paragrafo(p, extra: list[str]) -> str:
    """Il testo di un <w:p>: testo, tabulazioni e a capo; le caselle di testo dentro il paragrafo vanno in `extra` (righe a sé),
    il ramo di riserva dei contenuti alternativi (mc:Fallback) si salta (è un doppione della casella)."""
    parti: list[str] = []

    def giu(el):
        for ch in el:
            tag = ch.tag
            if not isinstance(tag, str):
                continue
            if tag == _MC_FALLBACK:
                continue
            if tag == _W + "txbxContent":
                extra.extend(_blocchi(ch))
                continue
            if tag == _W + "t":
                parti.append(ch.text or "")
            elif tag == _W + "tab":
                parti.append("\t")
            elif tag in (_W + "br", _W + "cr"):
                parti.append("\n")
            elif tag == _W + "noBreakHyphen":
                parti.append("-")
            elif tag in (_W + "delText", _W + "instrText"):
                continue          # testo cancellato in revisione, codici dei campi
            else:
                giu(ch)

    giu(p)
    return "".join(parti)


def _testo_tabella(tbl) -> list[str]:
    righe: list[str] = []
    for tr in tbl.findall(_W + "tr"):          # solo le righe di QUESTA tabella: le annidate passano dalla cella
        celle: list[str] = []
        for tc in tr:
            if tc.tag != _W + "tc":
                continue
            testo = " / ".join(x for x in (_testo_blocco_cella(tc)) if x.strip())
            celle.append(testo.strip())
        while celle and not celle[-1]:
            celle.pop()
        if any(c for c in celle):
            righe.append(" | ".join(celle))
    return righe


def _testo_blocco_cella(tc) -> list[str]:
    out: list[str] = []
    for ch in tc:
        if ch.tag == _W + "p":
            extra: list[str] = []
            t = _testo_paragrafo(ch, extra)
            if t.strip():
                out.append(t.strip())
            out.extend(e for e in extra if e.strip())
        elif ch.tag == _W + "tbl":
            out.extend(_testo_tabella(ch))
    return out


def _blocchi(contenitore) -> list[str]:
    """Il corpo (o un'intestazione, una casella di testo) NELL'ORDINE: paragrafi e tabelle, controlli di contenuto compresi."""
    righe: list[str] = []
    for ch in contenitore:
        tag = ch.tag
        if not isinstance(tag, str):
            continue
        if tag == _W + "p":
            extra: list[str] = []
            righe.append(_testo_paragrafo(ch, extra))
            righe.extend(extra)
        elif tag == _W + "tbl":
            righe.extend(_testo_tabella(ch))
        elif tag == _W + "sdt":                     # controllo di contenuto: il testo sta in w:sdtContent
            contenuto = ch.find(_W + "sdtContent")
            if contenuto is not None:
                righe.extend(_blocchi(contenuto))
        elif tag in (_W + "customXml", _W + "smartTag"):
            righe.extend(_blocchi(ch))
    return righe


def _note_docx(path: Path) -> list[str]:
    """Le note a piè di pagina e di chiusura (python-docx non le legge): «[n] testo»; mai i separatori."""
    out: list[str] = []
    try:
        from lxml import etree
    except Exception:  # noqa: BLE001
        return out
    try:
        with zipfile.ZipFile(path) as z:
            for nome, tag in (("word/footnotes.xml", "footnote"), ("word/endnotes.xml", "endnote")):
                if nome not in z.namelist():
                    continue
                radice = etree.fromstring(z.read(nome))
                for n in radice.iter(_W + tag):
                    tipo = n.get(_W + "type") or ""
                    if tipo in ("separator", "continuationSeparator", "continuationNotice"):
                        continue
                    testo = " ".join(r.strip() for r in _blocchi(n) if r.strip())
                    if testo:
                        out.append(f"[{n.get(_W + 'id') or '?'}] {testo}")
    except Exception:  # noqa: BLE001
        return out
    return out


def _read_docx(path: Path) -> ReadResult:
    try:
        from docx import Document
        doc = Document(str(path))
    except Exception as e:
        return ReadResult(text="", source_format="docx", ok=False, error=str(e))
    try:
        corpo = _blocchi(doc.element.body)
        testa, piede, visti = [], [], set()
        for sez in doc.sections:
            for parte, dove in ((getattr(sez, "header", None), testa), (getattr(sez, "first_page_header", None), testa),
                                (getattr(sez, "footer", None), piede), (getattr(sez, "first_page_footer", None), piede)):
                try:
                    if parte is None or getattr(parte, "is_linked_to_previous", False):
                        continue
                    righe = [r for r in _blocchi(parte._element) if r.strip()]
                except Exception:  # noqa: BLE001
                    continue
                chiave = "\n".join(righe)
                if righe and chiave not in visti:
                    visti.add(chiave)
                    dove.extend(righe)
        note = _note_docx(path)
        pezzi = []
        if testa:
            pezzi.append("\n".join(testa))
        pezzi.append("\n".join(corpo))
        if piede:
            pezzi.append("\n".join(piede))
        if note:
            pezzi.append("\n".join(note))
        testo = "\n\n".join(p.strip("\n") for p in pezzi if p.strip())
        testo = re.sub(r"\n{4,}", "\n\n\n", testo)
        return ReadResult(text=testo.strip(), source_format="docx")
    except Exception as e:  # noqa: BLE001 — un docx strano: almeno i paragrafi, come prima
        try:
            return ReadResult(text="\n".join(p.text for p in doc.paragraphs).strip(), source_format="docx")
        except Exception:  # noqa: BLE001
            return ReadResult(text="", source_format="docx", ok=False, error=str(e))


# ── OpenDocument (.odt) ───────────────────────────────────────────────────────

_ODT_TEXT = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
_ODT_TABLE = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
_ODT_OFFICE = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"


def _odt_inline(el) -> str:
    parti = [el.text or ""]
    for ch in el:
        tag = ch.tag if isinstance(ch.tag, str) else ""
        if tag == _ODT_TEXT + "s":
            parti.append(" " * int(ch.get(_ODT_TEXT + "c") or 1))
        elif tag == _ODT_TEXT + "tab":
            parti.append("\t")
        elif tag == _ODT_TEXT + "line-break":
            parti.append("\n")
        elif tag in (_ODT_TEXT + "note-citation",):
            pass
        else:
            parti.append(_odt_inline(ch))
        parti.append(ch.tail or "")
    return "".join(parti)


def _odt_blocchi(el) -> list[str]:
    righe: list[str] = []
    for ch in el:
        tag = ch.tag if isinstance(ch.tag, str) else ""
        if tag in (_ODT_TEXT + "p", _ODT_TEXT + "h"):
            righe.append(_odt_inline(ch))
        elif tag == _ODT_TABLE + "table":
            for riga in ch.iter(_ODT_TABLE + "table-row"):
                celle = [" / ".join(x.strip() for x in _odt_blocchi(c) if x.strip())
                         for c in riga if isinstance(c.tag, str) and c.tag == _ODT_TABLE + "table-cell"]
                if any(celle):
                    righe.append(" | ".join(celle))
        else:
            righe.extend(_odt_blocchi(ch))
    return righe


def _read_odt(path: Path) -> ReadResult:
    try:
        from lxml import etree
        with zipfile.ZipFile(path) as z:
            radice = etree.fromstring(z.read("content.xml"))
        corpo = radice.find(_ODT_OFFICE + "body")
        righe = _odt_blocchi(corpo if corpo is not None else radice)
        return ReadResult(text="\n".join(righe).strip(), source_format="odt")
    except Exception as e:  # noqa: BLE001
        return ReadResult(text="", source_format="odt", ok=False, error=str(e))


# ── RTF ───────────────────────────────────────────────────────────────────────

# destinazioni che NON sono testo del documento (tabelle dei font e dei colori, stili, immagini, metadati, campi nascosti)
_RTF_SALTA = {
    "fonttbl", "colortbl", "stylesheet", "info", "pict", "object", "themedata", "colorschememapping", "latentstyles",
    "datastore", "xmlnstbl", "listtable", "listoverridetable", "rsidtbl", "generator", "filetbl", "revtbl", "header",
    "footer", "headerl", "headerr", "headerf", "footerl", "footerr", "footerf", "fldinst", "expandedcolortbl", "mmathPr",
    "operator", "author", "title", "subject", "keywords", "company", "comment", "doccomm", "bkmkstart", "bkmkend",
    "pgdsctbl", "listtext", "pntext", "pntxtb", "pntxta", "wgrffmtfilter", "panose", "falt", "fcharset",
}
_RTF_TOKEN = re.compile(r"\\([a-zA-Z]{1,32})(-?\d{1,10})? ?|\\'([0-9a-fA-F]{2})|\\([^a-zA-Z])|([{}])|([^\\{}]+)",
                        re.S)


def rtf_a_testo(raw: str) -> str:
    """RTF → testo semplice: paragrafi e tabulazioni, \\'hh nella code page del documento, \\uN (Unicode) col numero di
    caratteri di riserva \\ucN, le destinazioni non testuali saltate. Pensato per i documenti legali, non per l'impaginazione."""
    cp = "cp1252"
    m = re.search(r"\\ansicpg(\d+)", raw[:2000])
    if m:
        cp = "cp" + m.group(1)
        try:
            "".encode(cp)
        except LookupError:
            cp = "cp1252"
    pila: list[tuple[bool, int]] = []
    salta, uc = False, 1
    da_saltare = 0          # caratteri di riserva dopo \uN
    out: list[str] = []
    asterisco = False
    for tok in _RTF_TOKEN.finditer(raw):
        parola, num, hexa, simbolo, graffa, testo = tok.groups()
        if graffa == "{":
            pila.append((salta, uc))
            asterisco = False
            continue
        if graffa == "}":
            if pila:
                salta, uc = pila.pop()
            asterisco = False
            continue
        if simbolo is not None:
            if simbolo == "*":
                asterisco = True          # destinazione opzionale: si salta se non la conosciamo
                salta = True
                continue
            if salta:
                continue
            if simbolo in ("\n", "\r"):
                out.append("\n")          # «\» a fine riga = nuovo paragrafo (TextEdit, Word)
            elif simbolo == "~":
                out.append("\u00a0")
            elif simbolo in ("-", "_"):
                out.append("-" if simbolo == "_" else "")
            elif simbolo in ("\\", "{", "}"):
                out.append(simbolo)
            continue
        if hexa is not None:
            if da_saltare:
                da_saltare -= 1
                continue
            if not salta:
                try:
                    out.append(bytes([int(hexa, 16)]).decode(cp))
                except Exception:  # noqa: BLE001
                    pass
            continue
        if parola is not None:
            if parola in _RTF_SALTA:
                salta = True
                continue
            if asterisco:
                salta = True
                continue
            if salta:
                continue
            if parola == "uc" and num:
                uc = max(0, int(num))
            elif parola == "u" and num:
                n = int(num)
                if n < 0:
                    n += 65536
                try:
                    out.append(chr(n))
                except ValueError:
                    pass
                da_saltare = uc
            elif parola in ("par", "line", "sect", "page"):
                out.append("\n")
            elif parola == "row":
                out.append("\n")
            elif parola == "cell":
                out.append(" | ")
            elif parola == "tab":
                out.append("\t")
            elif parola in ("emdash",):
                out.append("—")
            elif parola in ("endash",):
                out.append("–")
            elif parola in ("lquote", "rquote"):
                out.append("'")
            elif parola in ("ldblquote", "rdblquote"):
                out.append('"')
            elif parola == "bullet":
                out.append("•")
            continue
        if testo is not None:
            if salta:
                continue
            pezzo = testo.replace("\r", "").replace("\n", "")
            if da_saltare and pezzo:
                tolti = min(da_saltare, len(pezzo))
                pezzo = pezzo[tolti:]
                da_saltare -= tolti
            out.append(pezzo)
    testo = "".join(out)
    testo = re.sub(r"[ \t]+\n", "\n", testo)
    testo = re.sub(r"\n{3,}", "\n\n", testo)
    return testo.strip()


def _read_rtf(path: Path) -> ReadResult:
    try:
        raw = path.read_bytes().decode("latin-1")     # RTF è ASCII a 7 bit: i caratteri veri sono nei codici
        return ReadResult(text=rtf_a_testo(raw), source_format="rtf")
    except Exception as e:  # noqa: BLE001
        return ReadResult(text="", source_format="rtf", ok=False, error=str(e))


# ── HTML, testo ───────────────────────────────────────────────────────────────


def _read_html(path: Path) -> ReadResult:
    try:
        from selectolax.parser import HTMLParser
        raw = path.read_text(encoding="utf-8", errors="replace")
        tree = HTMLParser(raw)
        # HUDOC ships inline <style> blocks — drop them first.
        for tag in ("style", "script"):
            for node in tree.css(tag):
                node.decompose()
        text = tree.body.text(separator="\n", strip=True) if tree.body else tree.text(separator="\n", strip=True)
        return ReadResult(text=text.strip(), source_format="html")
    except Exception as e:
        return ReadResult(text="", source_format="html", ok=False, error=str(e))


def _read_plain(path: Path) -> ReadResult:
    try:
        return ReadResult(
            text=path.read_text(encoding="utf-8", errors="replace").strip(),
            source_format=path.suffix.lstrip(".").lower() or "txt",
        )
    except Exception as e:
        return ReadResult(text="", source_format="txt", ok=False, error=str(e))


_PER_FORMATO = {"docx": _read_docx, "doc": _read_doc, "odt": _read_odt, "rtf": _read_rtf, "pdf": _read_pdf}


def read_text(path: str | Path) -> ReadResult:
    p = Path(path)
    if not p.exists():
        return ReadResult(
            text="", source_format="missing", ok=False, error="file not found"
        )
    ext = p.suffix.lower()
    reale = formato_reale(p)
    lettore = _PER_FORMATO.get(reale or "")
    if lettore is not None:
        if reale != ext.lstrip("."):
            log.info("documento %s: estensione %s, formato reale %s", p.name, ext, reale)
        return lettore(p)
    if reale == "html" and ext not in {".txt"}:
        return _read_html(p)
    if ext == ".pdf":
        return _read_pdf(p)
    if ext == ".doc":
        return _read_doc(p)
    if ext == ".docx":
        return _read_docx(p)
    if ext == ".odt":
        return _read_odt(p)
    if ext in {".html", ".htm"}:
        return _read_html(p)
    if ext == ".rtf":
        return _read_plain(p)     # non comincia con {\rtf: è testo semplice con l'estensione sbagliata
    if ext == ".txt":
        return _read_plain(p)
    return ReadResult(
        text="", source_format=ext or "unknown", ok=False,
        error=f"unsupported extension: {ext}",
    )
